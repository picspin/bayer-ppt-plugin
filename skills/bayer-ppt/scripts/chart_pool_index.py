#!/usr/bin/env python3
"""List/search Chart Pool (cp) or BC Chart Library (bc) slides by keyword (title + all text + layout name).
Usage: chart_pool_index.py [--lib cp|bc] [--grep KW ...] [--show N] [--rebuild]
Index cached at ~/.cache/bayer-ppt/<lib>_index.json. Indices are 1-based file order."""
import argparse, json, os
from pptx import Presentation
SRC = {"cp": "/home/box/tools/bayer-ppt/source/official-template.pptx",
       "bc": "/home/box/tools/bayer-ppt/source/structured-graphics.pptx"}
def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6: yield from walk(sh.shapes)
        elif sh.has_text_frame and sh.text_frame.text.strip(): yield sh.text_frame.text.strip()
def build(lib):
    prs = Presentation(SRC[lib]); rows = []
    for i, s in enumerate(prs.slides, 1):
        t = list(walk(s.shapes)); title = s.shapes.title.text.strip() if s.shapes.title is not None and s.shapes.title.has_text_frame else (t[0] if t else "")
        rows.append(dict(n=i, layout=s.slide_layout.name, title=title.replace("\n", " ")[:90], shapes=len(s.shapes), text=" | ".join(t)[:600]))
    return rows
def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--lib", default="cp", choices=SRC)
    ap.add_argument("--grep", nargs="*"); ap.add_argument("--show", type=int); ap.add_argument("--rebuild", action="store_true")
    a = ap.parse_args(); cache = os.path.expanduser(f"~/.cache/bayer-ppt/{a.lib}_index.json")
    if a.rebuild or not os.path.exists(cache):
        os.makedirs(os.path.dirname(cache), exist_ok=True); json.dump(build(a.lib), open(cache, "w"), ensure_ascii=False)
    rows = json.load(open(cache))
    if a.show:
        print(json.dumps(rows[a.show - 1], ensure_ascii=False, indent=1)); return
    for r in rows:
        hay = (r["title"] + " " + r["text"] + " " + r["layout"]).lower()
        if not a.grep or all(k.lower() in hay for k in a.grep):
            print(f"{a.lib.upper()}{r['n']:>3}  [{r['layout'][:24]}]  {r['title']}  (shapes={r['shapes']})")
if __name__ == "__main__": main()
