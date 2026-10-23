#!/usr/bin/env python3
"""Lint a pptx slide face (NOT speaker notes): banned words, text load, visual presence,
information density (target 0.70-0.80 of the content safe area), cited-content source crops (SRC_*),
cover fit (COVER_IMG aspect/bounds) and pattern diversity (INTENT/PATTERN lines in notes).
Usage: lint_slide_face.py deck.pptx [--max-boxes 8] [--max-chars 220] [--max-words 90]
       [--dmin 0.70] [--dmax 0.80] [--strict] [--json]
Exit 1 if any banned word on slide face; with --strict also exit 1 on citation-crop violations."""
import sys, re, json, argparse
from pptx import Presentation
from pptx.util import Emu
BANNED = ["红线","合规","规避","违规","禁区","不得宣传","超说明书","off-label","off label",
          "red line","redline","compliance","compliant","avoid claim","Smart Protocol"]
def texts(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # group
            yield from texts(sh.shapes); continue
        if sh.has_text_frame and sh.text_frame.text.strip():
            yield sh, sh.text_frame.text
        if getattr(sh, "has_table", False) and sh.has_table:
            for r in sh.table.rows:
                for c in r.cells:
                    if c.text.strip(): yield sh, c.text
def has_pic(shapes):
    for sh in shapes:
        if sh.shape_type in (13,) or getattr(sh,"has_chart",False) and sh.has_chart: return True
        if sh.shape_type == 6 and has_pic(sh.shapes): return True
        if sh.is_placeholder and sh.placeholder_format.type == 18: return True
    return False
IN = 914400
CITE_RE = re.compile(r"arXiv|doi[:.]|10\.\d{4,}/|PMID|et al\.?|\[\d+(?:[,\u2013\-]\d+)*\]|\b(?:Fig\.?|Figure|Table)\s*\d|文献|引自|来源[:：]|Source[:：]", re.I)
SKIP_RE = re.compile(r"agenda|目录|议程|reference|参考文献|section|chapter|章节|thank|谢谢|final slide|closing|contact|Q&A", re.I)
def flat(shapes):
    for sh in shapes:
        if sh.shape_type == 6: yield from flat(sh.shapes)
        else: yield sh
def visible(sh):
    if sh.shape_type in (13,) or getattr(sh, "has_chart", False) and sh.has_chart or getattr(sh, "has_table", False) and sh.has_table: return True
    if sh.has_text_frame and sh.text_frame.text.strip(): return True
    try:
        if sh.fill.type is not None and sh.fill.type != 5: return True  # filled autoshape (5 = background)
    except Exception: pass
    if getattr(sh, "shape_type", None) == 9 or sh.__class__.__name__ == "Connector": return True  # lines/arrows
    return False
def text_extent(sh):
    """Estimated ink box of a text frame (not its frame): lines x font height, longest line x char width."""
    sizes = [r.font.size.pt for pg in sh.text_frame.paragraphs for r in pg.runs if r.font.size] or [14]
    pt = sum(sizes)/len(sizes); emu_pt = 12700; w_cap = sh.width
    lines = 0; longest = 0
    for pg in sh.text_frame.paragraphs:
        t = "".join(r.text for r in pg.runs) or pg.text
        ln = sum(1.0 if ord(c) > 0x2e80 else 0.55 for c in t) * pt * emu_pt  # CJK ~1em, latin ~0.55em
        wraps = max(1, -(-int(ln) // max(1, int(w_cap)))) if ln else 1
        lines += wraps; longest = max(longest, min(ln, w_cap))
    return sh.left, sh.top, int(longest) + int(0.1*IN), min(sh.height, int(lines*pt*1.25*emu_pt) + int(0.1*IN))
def density(slide, W, H, gx=40, gy=24):
    """Proxy: union of content 'ink' boxes clipped to the safe area (0.5in sides, 1.3in title band, 0.6in footer)
    on a 40x24 grid. Pictures/charts/tables = bounds; text = estimated text extent; filled autoshapes count only
    when small (<12% of safe area: chips, icons, arrows). Large panels/backgrounds are containers, not content."""
    l, r, t, b = int(0.5*IN), W-int(0.5*IN), int(1.3*IN), H-int(0.6*IN)
    cw, ch = (r-l)/gx, (b-t)/gy; grid = [[False]*gx for _ in range(gy)]; sa = (r-l)*(b-t)
    title_ids = {id(slide.shapes.title._element)} if slide.shapes.title is not None else set()
    for sh in flat(slide.shapes):
        if sh.left is None or sh.width is None or not visible(sh): continue
        txt = sh.text_frame.text if sh.has_text_frame else ""
        if id(sh._element) in title_ids or re.match(r"(Title|Subtitle|Footer|Slide Number|Date)", sh.name) or "Bayer Radiology ·" in txt: continue
        if sh.is_placeholder and sh.placeholder_format.type in (1, 3, 4, 13, 15, 16): continue  # title, ctr title, subtitle, sld num, footer, date
        if sh.width >= W*0.95 and sh.height >= H*0.9: continue  # full-slide background
        pic = sh.shape_type == 13 or getattr(sh, "has_chart", False) and sh.has_chart or getattr(sh, "has_table", False) and sh.has_table
        if pic: x, y, w_, h_ = sh.left, sh.top, sh.width, sh.height
        elif txt.strip() and sh.width*sh.height >= 0.12*sa: x, y, w_, h_ = text_extent(sh)
        elif sh.width*sh.height >= 0.12*sa: continue
        else: x, y, w_, h_ = sh.left, sh.top, sh.width, sh.height
        x0, x1, y0, y1 = max(x, l), min(x+w_, r), max(y, t), min(y+h_, b)
        if x1 <= x0 or y1 <= y0: continue
        for j in range(int((y0-t)//ch), min(gy, int(-(-(y1-t)//ch)))):
            for i in range(int((x0-l)//cw), min(gx, int(-(-(x1-l)//cw)))): grid[j][i] = True
    return sum(map(sum, grid))/(gx*gy)
def src_crops(slide):
    return [sh.name for sh in flat(slide.shapes) if sh.name.upper().startswith("SRC_") and (sh.shape_type == 13 or (sh.is_placeholder and sh.placeholder_format.type == 18))]
def cover_check(prs, s):
    pics = [sh for sh in s.shapes if sh.shape_type == 13 or (sh.is_placeholder and sh.placeholder_format.type == 18)]
    if not pics: return ["cover: no image (see cover-and-density.md)"]
    big = max(pics, key=lambda p: p.width*p.height); w = []
    if big.name != "COVER_IMG": w.append(f"cover: main image '{big.name}' not COVER_IMG (use fit_cover_image.py)")
    try:
        c = big.crop_left+big.crop_right, big.crop_top+big.crop_bottom
        iw, ih = big.image.size; eff = (iw*(1-c[0]))/(ih*(1-c[1])); box = big.width/big.height
        if abs(eff/box-1) > 0.02: w.append(f"cover: image aspect {eff:.3f} != box {box:.3f} (stretched)")
    except Exception: pass
    if sum(1 for p in pics if p.width*p.height > 0.15*prs.slide_width*prs.slide_height) > 1: w.append("cover: multiple large images (tiled?)")
    return w
def main():
    ap = argparse.ArgumentParser(); ap.add_argument("pptx")
    ap.add_argument("--max-boxes", type=int, default=8); ap.add_argument("--max-chars", type=int, default=220)
    ap.add_argument("--max-words", type=int, default=90); ap.add_argument("--json", action="store_true")
    ap.add_argument("--dmin", type=float, default=0.70); ap.add_argument("--dmax", type=float, default=0.80)
    ap.add_argument("--strict", action="store_true")
    a = ap.parse_args(); prs = Presentation(a.pptx); out = []; hits = 0; cites = 0; n_sl = len(prs.slides); patterns = []
    for i, s in enumerate(prs.slides, 1):
        boxes = list(texts(s.shapes)); allt = "\n".join(t for _, t in boxes)
        cjk = len(re.findall(r"[\u4e00-\u9fff]", allt)); words = len(re.findall(r"[A-Za-z0-9][A-Za-z0-9.\-/%]*", allt))
        found = sorted({b for b in BANNED if b.lower() in allt.lower()})
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        w = []
        if found: w.append("BANNED:" + ",".join(found)); hits += 1
        nbox = len({id(sh) for sh, t in boxes if len(t.strip()) >= 12 and not re.fullmatch(r"[\d\s\-/.:]+", t.strip()) and "Bayer Radiology ·" not in t})  # substantive boxes; ignores labels, dates, footer
        if nbox > a.max_boxes: w.append(f"boxes {nbox}>{a.max_boxes}")
        if cjk > a.max_chars: w.append(f"cjk {cjk}>{a.max_chars}")
        if words > a.max_words: w.append(f"words {words}>{a.max_words}")
        if not has_pic(s.shapes): w.append("no image/chart")
        title = s.shapes.title.text_frame.text if s.shapes.title is not None and s.shapes.title.has_text_frame else ""
        skip = i == 1 or i == n_sl and SKIP_RE.search(title + s.slide_layout.name) or bool(SKIP_RE.search(title + " " + s.slide_layout.name)) or "DENSITY: skip" in notes
        d = None if skip else round(density(s, prs.slide_width, prs.slide_height), 3)
        if d is not None and d < a.dmin: w.append(f"density {d:.2f}<{a.dmin}")
        if d is not None and d > a.dmax: w.append(f"density {d:.2f}>{a.dmax}")
        cited = bool(CITE_RE.search(allt)) and not re.search(r"reference|参考文献", title, re.I); crops = src_crops(s)
        if cited and not crops and i != 1: w.append("CITE: cited content without SRC_ source crop"); cites += 1
        if i == 1: w += cover_check(prs, s)
        m = re.search(r"PATTERN:\s*([A-Za-z]+[\w-]*?)(?:\d|\b)", notes)
        if m and not skip: patterns.append(m.group(1).lower())
        out.append(dict(slide=i, boxes=nbox, cjk=cjk, words=words, banned=found, notes_chars=len(notes), density=d, cited=cited, src_crops=crops, warnings=w))
    content = [r for r in out if r["density"] is not None]
    div = []
    if patterns:
        from collections import Counter
        top, cnt = Counter(patterns).most_common(1)[0]
        if cnt > 0.3*len(content): div.append(f"pattern '{top}' on {cnt}/{len(content)} content slides (>30%)")
    elif content: div.append("no 'PATTERN:' lines in notes (chart-intent-map.md step 5)")
    if a.json: print(json.dumps(dict(slides=out, diversity=div), ensure_ascii=False, indent=1))
    else:
        for r in out:
            ds = "  -  " if r["density"] is None else f"{r['density']:.2f}"
            print(f"S{r['slide']:02d} boxes={r['boxes']:2d} cjk={r['cjk']:4d} words={r['words']:3d} notes={r['notes_chars']:4d} dens={ds} " + ("; ".join(r['warnings']) or "OK"))
        ins = sum(1 for r in content if a.dmin <= r["density"] <= a.dmax)
        avg = sum(r["density"] for r in content)/len(content) if content else 0
        print(f"DENSITY content_slides={len(content)} in_band={ins} low={sum(1 for r in content if r['density']<a.dmin)} high={sum(1 for r in content if r['density']>a.dmax)} mean={avg:.2f}")
        print("DIVERSITY " + ("; ".join(div) or "OK"))
        print(f"SUMMARY slides={len(out)} banned_slides={hits} cite_crop_violations={cites} warn_slides={sum(1 for r in out if r['warnings'])}")
    sys.exit(1 if hits or (a.strict and cites) else 0)
if __name__ == "__main__": main()
